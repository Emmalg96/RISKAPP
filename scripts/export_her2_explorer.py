#!/usr/bin/env python3
"""Export Explorer sheet from HER2 traceability workbook to JSON."""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_INPUT = Path("data/source/HER2_Risk_Traceability_PROJECT_AWARE_DRAFT.xlsx")
FALLBACK_INPUT = Path("/cursor/stores/self/data/HER2_Risk_Traceability_PROJECT_AWARE_DRAFT.xlsx")
DEFAULT_OUTPUT = Path("public/data/her2-explorer.json")
SHEET_NAME = "Explorer"
EXPECTED_HEADERS = {
    "Project",
    "RiskScope",
    "RiskSource",
    "RiskLine",
    "RiskSummary",
    "ProductRequirement",
    "Protocol",
}


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def export_sheet(input_path: Path, output_path: Path) -> dict[str, Any]:
    wb = load_workbook(input_path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet {SHEET_NAME!r} not found. Available: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise ValueError("Explorer sheet has no header row")
    headers = [cell_str(h) for h in header_row]
    while headers and not headers[-1]:
        headers.pop()

    missing = EXPECTED_HEADERS - set(headers)
    if missing:
        raise ValueError(f"Explorer sheet missing expected columns: {sorted(missing)}")

    records: list[dict[str, str]] = []
    for raw in rows_iter:
        if not any(raw):
            continue
        values = list(raw[: len(headers)])
        if len(values) < len(headers):
            values.extend([None] * (len(headers) - len(values)))
        record = {headers[i]: cell_str(values[i]) for i in range(len(headers)) if headers[i]}
        if not any(record.values()):
            continue
        records.append(record)

    payload = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceFile": str(input_path),
        "sheetName": SHEET_NAME,
        "rowCount": len(records),
        "rows": records,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return payload


def resolve_input() -> Path:
    if len(sys.argv) >= 2:
        return Path(sys.argv[1]).expanduser().resolve()
    if DEFAULT_INPUT.is_file():
        return DEFAULT_INPUT.resolve()
    if FALLBACK_INPUT.is_file():
        return FALLBACK_INPUT.resolve()
    raise FileNotFoundError(f"Workbook not found at {DEFAULT_INPUT} or {FALLBACK_INPUT}")


def main() -> int:
    try:
        input_path = resolve_input()
        output_path = Path(sys.argv[2]).expanduser().resolve() if len(sys.argv) >= 3 else DEFAULT_OUTPUT.resolve()
        payload = export_sheet(input_path, output_path)
        print(json.dumps({"ok": True, "output": str(output_path), "rows": payload["rowCount"]}, indent=2))
        return 0
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": str(exc)}), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
